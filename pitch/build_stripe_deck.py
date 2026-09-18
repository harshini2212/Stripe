"""Swap the Mercury job-pitch deck to Stripe / Stripify: text runs, notes, and images."""
import sys
from pptx import Presentation

src, out, pitch = sys.argv[1], sys.argv[2], sys.argv[3]
RULES = [
    ("Agentic Mercury", "Stripify"),
    ("modal-ai.up.railway.app", "stripify.up.railway.app"),
    ("Mercury Treasury", "Stripe Treasury"),
    ("Mercury Card", "Stripe Issuing"),
    ("Mercury Assistant", "Stripe Assistant"),
    ("MERCURY", "STRIPE"),
    ("Mercury", "Stripe"),
    ("Card · Stripe Treasury · AP", "Issuing · Treasury · AP"),
]
def fix(s):
    for a, b in RULES: s = s.replace(a, b)
    return s

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh

def swap_image(pic, path):
    rId = pic._element.blip_rId
    part = pic.part.related_part(rId)
    part._blob = open(path, "rb").read()

NOTES4 = """Architecture - the trust layer under Stripe's finance agents.
- Deterministic engine runs FIRST, for $0: fraud ensemble (Isolation Forest + gradient boosting + graph features, ROC-AUC ~0.95), policy replay (re-run a spend policy over history for its dollar impact), tie-out reconciliation (every expense report rolls up to the cent), duplicate-charge / redundant-subscription / limit detection. Each emits a Finding with evidence and dollars at risk.
- Where the numbers come from - real models: a fraud-ring graph (connected components over shared-device / cross-metro-IP edges), causal runway (do-operator on a structural causal model, not a re-plotted trend), a gradient-boosted PD model that sizes the Stripe Capital limit, and an idle-cash yield optimizer for the Stripe Treasury sweep.
- Only then does the LLM layer run: five narrow agents (categorize, policy audit, fraud triage, dispute adjudication, tie-out) sequenced by an orchestrator, escalating to a full fraud-investigator workflow. Opus 4.8 reviews, Sonnet 4.6 investigates, Haiku 4.5 parses receipts and docs (model routing). Every finding is tagged engine / llm / both.
- Every agent decision is graded against held-out ground truth on the financial-correctness benchmark (bootstrap CIs, cost, latency) - the promotion gate before a model ships. Anything risky (fraud over threshold, PII, an irreversible write, a tie-out mismatch, a benchmark below the gate) HALTS and asks a human to confirm.
- It is one orchestrated run: categorize -> policy -> fraud -> dispute -> tie-out, with the agent confirming before any write tool. That is the trust layer that lets Stripe act on real spend."""

p = Presentation(src)
n_runs = 0
for i, slide in enumerate(p.slides, 1):
    for sh in walk(slide.shapes):
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    new = fix(r.text)
                    if new != r.text: r.text = new; n_runs += 1
        if sh.shape_type == 13:
            if i == 1 and sh.shape_id == 58: swap_image(sh, f"{pitch}/stripify_qr.png")
            if i == 3 and sh.shape_id == 101: swap_image(sh, f"{pitch}/stripify_dash.png")
            if i == 4 and sh.shape_id == 2: swap_image(sh, f"{pitch}/stripify_arch.png")
            if i == 12 and sh.shape_id in (257, 258): swap_image(sh, f"{pitch}/stripify_qr.png")
    if slide.has_notes_slide:
        tf = slide.notes_slide.notes_text_frame
        if i == 4:
            tf.text = NOTES4
        elif tf.text.strip():
            for para in tf.paragraphs:
                for r in para.runs: r.text = fix(r.text)
p.save(out)
print("runs changed:", n_runs, "->", out)

# verify
p2 = Presentation(out)
txt = " ".join(r.text for s in p2.slides for sh in walk(s.shapes) if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs)
print("mercury left:", "mercury" in txt.lower(), "| stripify:", "stripify" in txt.lower(), "| link:", "stripify.up.railway.app" in txt)
